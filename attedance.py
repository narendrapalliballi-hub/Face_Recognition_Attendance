import cv2
import csv
import time
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook

# ---------------- CONFIGURATION ----------------

PROJECT_DIR = Path(__file__).resolve().parent
MODEL_PATH = PROJECT_DIR / "trainer" / "face_model.yml"
CASCADE_PATH = PROJECT_DIR / "haarcascade_frontalface_default.xml"
STUDENTS_FILE = PROJECT_DIR / "students.csv"
ATTENDANCE_FILE = PROJECT_DIR / "attendance.xlsx"

# LBPH: lower confidence value generally means a closer match.
# 85 is an initial test value; tune using actual test results.
CONFIDENCE_THRESHOLD = 85


# ---------------- LOAD STUDENTS ----------------

def load_students():
    students = {}

    if not STUDENTS_FILE.exists():
        raise FileNotFoundError(
            f"Student file not found: {STUDENTS_FILE}"
        )

    with open(
        STUDENTS_FILE, "r",
        newline="", encoding="utf-8-sig"
    ) as file:
        reader = csv.DictReader(file)

        if not reader.fieldnames:
            raise ValueError("students.csv is empty.")

        # Match column names without depending on capitalization.
        columns = {
            col.strip().lower(): col
            for col in reader.fieldnames if col
        }

        if "student_id" not in columns or "student_name" not in columns:
            raise ValueError(
                "CSV headers must be Student_ID,Student_Name"
            )

        for row in reader:
            sid = (row.get(columns["student_id"]) or "").strip()
            name = (row.get(columns["student_name"]) or "").strip()

            if sid and name:
                students[sid] = name

    return students


# ---------------- MAIN PROGRAM ----------------

def main():
    for path in (MODEL_PATH, CASCADE_PATH, STUDENTS_FILE):
        if not path.exists():
            print(f"ERROR: Missing file: {path}")
            return

    if not hasattr(cv2, "face"):
        print("ERROR: OpenCV face module is missing.")
        print("Run: python -m pip install opencv-contrib-python")
        return

    try:
        students = load_students()
    except (OSError, ValueError) as error:
        print("Student loading error:", error)
        return

    if not students:
        print("ERROR: No students found in students.csv.")
        return

    print("Loaded students:", students)

    detector = cv2.CascadeClassifier(str(CASCADE_PATH))
    if detector.empty():
        print("ERROR: Could not load face cascade.")
        return

    recognizer = cv2.face.LBPHFaceRecognizer_create()

    try:
        recognizer.read(str(MODEL_PATH))
    except cv2.error as error:
        print("Could not load face model:", error)
        return

    # Create Excel workbook if it does not exist.
    if not ATTENDANCE_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Attendance"
        sheet.append([
            "Student ID", "Student Name",
            "Date", "Time", "Status"
        ])
        workbook.save(ATTENDANCE_FILE)
        workbook.close()

    try:
        workbook = load_workbook(ATTENDANCE_FILE)
    except Exception as error:
        print("Cannot open attendance.xlsx:", error)
        print("Close Excel if the file is open, then retry.")
        return

    if "Attendance" not in workbook.sheetnames:
        sheet = workbook.create_sheet("Attendance")
        sheet.append([
            "Student ID", "Student Name",
            "Date", "Time", "Status"
        ])
    else:
        sheet = workbook["Attendance"]

    # Prevent duplicate attendance for the same day.
    today = datetime.now().strftime("%Y-%m-%d")
    marked_today = set()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if (
            row[0] is not None
            and row[2] is not None
            and str(row[2])[:10] == today
        ):
            marked_today.add(str(row[0]).strip())

    camera = None

    try:
        # Open webcam.
        camera = cv2.VideoCapture(0, cv2.CAP_DSHOW)

        if not camera.isOpened():
            camera.release()
            camera = cv2.VideoCapture(0)

        if not camera.isOpened():
            print("ERROR: Webcam could not be opened.")
            return

        camera.set(cv2.CAP_PROP_FRAME_WIDTH, 640)
        camera.set(cv2.CAP_PROP_FRAME_HEIGHT, 480)

        print("\nFace attendance started.")
        print("Click webcam window and press Q to exit.")
        print("Watch the terminal for ID and confidence values.")

        last_debug = None

        while True:
            success, frame = camera.read()

            if not success:
                print("Could not read webcam frame.")
                break

            gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)

            faces = detector.detectMultiScale(
                gray,
                scaleFactor=1.2,
                minNeighbors=5,
                minSize=(80, 80)
            )

            for (x, y, w, h) in faces:
                face = gray[y:y + h, x:x + w]
                face = cv2.resize(face, (200, 200))

                student_id, confidence = recognizer.predict(face)
                sid = str(student_id)

                # Print only when prediction changes.
                debug_info = (
                    sid,
                    round(float(confidence), 2)
                )

                if debug_info != last_debug:
                    print(
                        f"Predicted ID: {sid}, "
                        f"Confidence: {confidence:.2f}, "
                        f"ID exists in CSV: {sid in students}"
                    )
                    last_debug = debug_info

                if (
                    sid in students
                    and confidence < CONFIDENCE_THRESHOLD
                ):
                    name = students[sid]
                    color = (0, 255, 0)
                    label = name

                    if sid not in marked_today:
                        now = datetime.now()

                        sheet.append([
                            sid,
                            name,
                            now.strftime("%Y-%m-%d"),
                            now.strftime("%H:%M:%S"),
                            "Present"
                        ])

                        workbook.save(ATTENDANCE_FILE)
                        marked_today.add(sid)
                        print(f"Attendance marked for {name}")

                else:
                    color = (0, 0, 255)
                    label = "Unknown"

                cv2.rectangle(
                    frame, (x, y), (x + w, y + h),
                    color, 2
                )

                cv2.putText(
                    frame,
                    label,
                    (x, max(y - 10, 25)),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.7,
                    color,
                    2
                )

                cv2.putText(
                    frame,
                    f"Confidence: {confidence:.1f}",
                    (x, y + h + 25),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    0.55,
                    color,
                    2
                )

            cv2.imshow(
                "Face Attendance - Press Q to quit",
                frame
            )

            key = cv2.waitKey(10) & 0xFF

            if key in (ord("q"), ord("Q"), 27):
                print("Closing webcam...")
                break

    except KeyboardInterrupt:
        print("\nStopped by user.")

    except Exception as error:
        print("Program error:", error)

    finally:
        if camera is not None:
            camera.release()

        cv2.destroyAllWindows()
        cv2.waitKey(1)

        workbook.close()
        print("Webcam closed.")
        print(f"Attendance file: {ATTENDANCE_FILE}")


if __name__ == "__main__":
    main()